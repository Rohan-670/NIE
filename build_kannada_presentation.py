import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# 16:9 Widescreen Presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color Palette (Karnataka Theme & Modern Dark Healthcare/Tech)
COLOR_BG = RGBColor(11, 19, 43)          # Deep Midnight Slate (#0B132B)
COLOR_CARD = RGBColor(22, 33, 62)        # Royal Deep Blue/Slate Card (#16213E)
COLOR_CARD_ALT = RGBColor(15, 23, 42)    # Dark Slate (#0F172A)

COLOR_KRED = RGBColor(220, 38, 38)       # Karnataka Flag Red (#DC2626)
COLOR_KGOLD = RGBColor(245, 158, 11)     # Karnataka Turmeric Gold (#F59E0B)
COLOR_KGOLD_LIGHT = RGBColor(254, 240, 138) # Light Golden (#FEF08A)

COLOR_TEAL = RGBColor(20, 184, 166)      # Modern Healthcare Teal (#14B8A6)
COLOR_TEAL_LIGHT = RGBColor(94, 234, 212)# Bright Teal (#5EEAD4)

COLOR_GREEN = RGBColor(34, 197, 94)      # Success Green (#22C55E)
COLOR_WHITE = RGBColor(255, 255, 255)    # Pure White
COLOR_MUTED = RGBColor(203, 213, 225)    # Light Slate Gray (#CBD5E1)
COLOR_BORDER = RGBColor(51, 65, 85)      # Card Border (#334155)

FONT_NAME = "Nirmala UI"

def set_slide_background(slide, color=COLOR_BG):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_karnataka_ribbon(slide):
    # Top dual-tone Karnataka Flag Ribbon (Red & Yellow)
    h = Inches(0.12)
    w_half = prs.slide_width / 2

    # Red bar (Left)
    red_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), w_half, h)
    red_bar.fill.solid()
    red_bar.fill.fore_color.rgb = COLOR_KRED
    red_bar.line.fill.background()

    # Yellow/Gold bar (Right)
    gold_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, w_half, Inches(0), w_half, h)
    gold_bar.fill.solid()
    gold_bar.fill.fore_color.rgb = COLOR_KGOLD
    gold_bar.line.fill.background()

def add_footer(slide, current_slide, total_slides=7):
    # Footer ribbon line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(40, 50, 75)
    line.line.fill.background()

    # Footer text
    tx = slide.shapes.add_textbox(Inches(0.8), Inches(7.1), Inches(6.0), Inches(0.35))
    tf = tx.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "ಹೀಲಿಯೋ (HEALIO) — ಕನ್ನಡ ದೈನಂದಿನ ಜೀವನದಲ್ಲಿ • ಕೈಮುಕ್ತ ಆಸ್ಪತ್ರೆ ಸೇವೆ"
    p.font.name = FONT_NAME
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_MUTED

    # Slogan in center/right
    tx2 = slide.shapes.add_textbox(Inches(6.8), Inches(7.1), Inches(4.5), Inches(0.35))
    tf2 = tx2.text_frame
    p2 = tf2.paragraphs[0]
    p2.text = "👑 ಸಿರಿಗನ್ನಡಂ ಗೆಲ್ಗೆ, ಸಿರಿಗನ್ನಡಂ ಬಾಳ್ಗೆ! 👑"
    p2.font.name = FONT_NAME
    p2.font.size = Pt(11)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_KGOLD

    # Page number
    tx3 = slide.shapes.add_textbox(Inches(11.5), Inches(7.1), Inches(1.0), Inches(0.35))
    tf3 = tx3.text_frame
    p3 = tf3.paragraphs[0]
    p3.alignment = PP_ALIGN.RIGHT
    p3.text = f"ಪುಟ {current_slide} / {total_slides}"
    p3.font.name = FONT_NAME
    p3.font.size = Pt(11)
    p3.font.color.rgb = COLOR_MUTED

def add_header(slide, title, subtitle, badge="ಕರ್ನಾಟಕ ನಾವೀನ್ಯತೆ • ೧೦೦% ಶುದ್ಧ ಕನ್ನಡ"):
    # Header Badge Pill
    if badge:
        badge_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.35), Inches(4.2), Inches(0.36))
        badge_box.fill.solid()
        badge_box.fill.fore_color.rgb = RGBColor(30, 41, 59)
        badge_box.line.color.rgb = COLOR_KGOLD
        badge_box.line.width = Pt(1.5)
        tf_b = badge_box.text_frame
        p_b = tf_b.paragraphs[0]
        p_b.alignment = PP_ALIGN.CENTER
        p_b.text = f"✨ {badge}"
        p_b.font.name = FONT_NAME
        p_b.font.size = Pt(12)
        p_b.font.bold = True
        p_b.font.color.rgb = COLOR_KGOLD_LIGHT

    # Main Title
    y_title = Inches(0.78) if badge else Inches(0.45)
    tx_title = slide.shapes.add_textbox(Inches(0.8), y_title, Inches(11.733), Inches(0.65))
    tf_title = tx_title.text_frame
    p_title = tf_title.paragraphs[0]
    p_title.text = title
    p_title.font.name = FONT_NAME
    p_title.font.size = Pt(28)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_WHITE

    # Subtitle
    if subtitle:
        tx_sub = slide.shapes.add_textbox(Inches(0.8), y_title + Inches(0.55), Inches(11.733), Inches(0.45))
        tf_sub = tx_sub.text_frame
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = subtitle
        p_sub.font.name = FONT_NAME
        p_sub.font.size = Pt(14)
        p_sub.font.color.rgb = COLOR_TEAL_LIGHT

# ==============================================================================
# SLIDE 1: ಮುಖಪುಟ (Title Slide)
# ==============================================================================
slide1 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(slide1)
add_karnataka_ribbon(slide1)
add_footer(slide1, 1)

# Large Center Decorative Card
card1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.1), Inches(10.933), Inches(5.6))
card1.fill.solid()
card1.fill.fore_color.rgb = COLOR_CARD
card1.line.color.rgb = COLOR_KGOLD
card1.line.width = Pt(2.5)

# Badge Pill
pill1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.3), Inches(1.4), Inches(4.7), Inches(0.45))
pill1.fill.solid()
pill1.fill.fore_color.rgb = RGBColor(185, 28, 28)
pill1.line.color.rgb = COLOR_KGOLD
pill1.line.width = Pt(1.5)
tf_p1 = pill1.text_frame
p_p1 = tf_p1.paragraphs[0]
p_p1.alignment = PP_ALIGN.CENTER
p_p1.text = "✨ ಕರ್ನಾಟಕ ರಾಜ್ಯ ನಾವೀನ್ಯತೆ • ೧೦೦% ಕನ್ನಡ ಆವೃತ್ತಿ ✨"
p_p1.font.name = FONT_NAME
p_p1.font.size = Pt(13)
p_p1.font.bold = True
p_p1.font.color.rgb = COLOR_KGOLD_LIGHT

# Big Hero Title
tx1_t = slide1.shapes.add_textbox(Inches(1.5), Inches(2.0), Inches(10.333), Inches(1.0))
tf1_t = tx1_t.text_frame
p1_t = tf1_t.paragraphs[0]
p1_t.alignment = PP_ALIGN.CENTER
p1_t.text = "ಹೀಲಿಯೋ (HEALIO)"
p1_t.font.name = FONT_NAME
p1_t.font.size = Pt(46)
p1_t.font.bold = True
p1_t.font.color.rgb = COLOR_WHITE

# Subtitle
tx1_s = slide1.shapes.add_textbox(Inches(1.5), Inches(2.95), Inches(10.333), Inches(0.8))
tf1_s = tx1_s.text_frame
p1_s = tf1_s.paragraphs[0]
p1_s.alignment = PP_ALIGN.CENTER
p1_s.text = "ಕನ್ನಡ ದೈನಂದಿನ ಜೀವನದಲ್ಲಿ — ಕೈಮುಕ್ತ ಆಸ್ಪತ್ರೆ ಸೇವೆ ಮತ್ತು ನೈಜ ಕನ್ನಡ ಕಲಿಕಾ ವ್ಯವಸ್ಥೆ"
p1_s.font.name = FONT_NAME
p1_s.font.size = Pt(20)
p1_s.font.bold = True
p1_s.font.color.rgb = COLOR_KGOLD

# Vision description
tx1_d = slide1.shapes.add_textbox(Inches(1.8), Inches(3.75), Inches(9.733), Inches(0.9))
tf1_d = tx1_d.text_frame
tf1_d.word_wrap = True
p1_d = tf1_d.paragraphs[0]
p1_d.alignment = PP_ALIGN.CENTER
p1_d.text = "ಆರೋಗ್ಯ, ಕೃತಕ ಬುದ್ಧಿಮತ್ತೆ (AI) ಮತ್ತು ಕನ್ನಡ ಭಾಷೆಯ ಸಾರ್ಥಕ ಸಂಗಮ.\nಯಾವುದೇ ಟೈಪಿಂಗ್ ಇಲ್ಲದೆ ಮಾತು, ಕಣ್ಣು ನೋಟ ಹಾಗೂ ಸ್ಪರ್ಶದ ಮೂಲಕ ಪ್ರತಿಯೊಬ್ಬ ಕನ್ನಡಿಗನಿಗೂ ಸಮಾನ ಆರೋಗ್ಯ ಸೌಲಭ್ಯ."
p1_d.font.name = FONT_NAME
p1_d.font.size = Pt(15)
p1_d.font.color.rgb = COLOR_MUTED

# 3 Pillars row
p_data = [
    ("🎙️ ಧ್ವನಿ ಆಧಾರಿತ ಬುಕಿಂಗ್", "ಕನ್ನಡ ಮಾತಿನಿಂದ ನೇರ ಬುಕಿಂಗ್"),
    ("👁️ ಕಣ್ಣು ನೋಟ ವಿಧಾನ", "೧.೨ ಸೆಕೆಂಡ್ SVG Dwell ಟೈಮರ್"),
    ("🎓 ದೈನಂದಿನ ಕನ್ನಡ ಸಂಗಾತಿ", "ಪ್ರಯಾಣ, ಔಷಧಾಲಯ & ಕೆಲಸ")
]
for i, (p_t, p_st) in enumerate(p_data):
    box = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.8 + i * 3.35), Inches(4.8), Inches(3.05), Inches(1.2))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(15, 23, 42)
    box.line.color.rgb = COLOR_TEAL
    box.line.width = Pt(1.5)
    tf_b = box.text_frame
    p_b1 = tf_b.paragraphs[0]
    p_b1.alignment = PP_ALIGN.CENTER
    p_b1.text = p_t
    p_b1.font.name = FONT_NAME
    p_b1.font.size = Pt(14)
    p_b1.font.bold = True
    p_b1.font.color.rgb = COLOR_TEAL_LIGHT
    p_b2 = tf_b.add_paragraph()
    p_b2.alignment = PP_ALIGN.CENTER
    p_b2.text = p_st
    p_b2.font.name = FONT_NAME
    p_b2.font.size = Pt(12)
    p_b2.font.color.rgb = COLOR_MUTED

# Slogan at bottom of card
tx1_slg = slide1.shapes.add_textbox(Inches(2.0), Inches(6.15), Inches(9.333), Inches(0.45))
tf1_slg = tx1_slg.text_frame
p1_slg = tf1_slg.paragraphs[0]
p1_slg.alignment = PP_ALIGN.CENTER
p1_slg.text = "👑 ಸಿರಿಗನ್ನಡಂ ಗೆಲ್ಗೆ, ಸಿರಿಗನ್ನಡಂ ಬಾಳ್ಗೆ! 👑"
p1_slg.font.name = FONT_NAME
p1_slg.font.size = Pt(15)
p1_slg.font.bold = True
p1_slg.font.color.rgb = COLOR_KGOLD

# ==============================================================================
# SLIDE 2: ಸಮಸ್ಯೆಯ ಹಿನ್ನೆಲೆ & ನವೀನ ಪರಿಹಾರ
# ==============================================================================
slide2 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(slide2)
add_karnataka_ribbon(slide2)
add_header(slide2, "ಸಮಸ್ಯೆಯ ಹಿನ್ನೆಲೆ ಮತ್ತು ನವೀನ ಪರಿಹಾರ", "ಸಾರ್ವಜನಿಕ ಆರೋಗ್ಯ ಕೇಂದ್ರಗಳಲ್ಲಿ ಭಾಷೆ ಮತ್ತು ಪ್ರವೇಶಿಸುವಿಕೆಯ ಸವಾಲುಗಳಿಗೆ ತಂತ್ರಜ್ಞಾನದ ಉತ್ತರ")
add_footer(slide2, 2)

# Left Card: ಸವಾಲುಗಳು (The Challenges)
card2_l = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.85), Inches(5.65), Inches(4.95))
card2_l.fill.solid()
card2_l.fill.fore_color.rgb = COLOR_CARD
card2_l.line.color.rgb = COLOR_KRED
card2_l.line.width = Pt(2.0)

tx2_lh = slide2.shapes.add_textbox(Inches(1.1), Inches(2.05), Inches(5.0), Inches(0.55))
p2_lh = tx2_lh.text_frame.paragraphs[0]
p2_lh.text = "🚨 ಪ್ರಸ್ತುತ ಸವಾಲುಗಳು (The Challenges)"
p2_lh.font.name = FONT_NAME
p2_lh.font.size = Pt(20)
p2_lh.font.bold = True
p2_lh.font.color.rgb = COLOR_KRED

challenges = [
    ("೧. ಆಸ್ಪತ್ರೆಗಳಲ್ಲಿ ದೀರ್ಘ ಸರತಿ ಸಾಲು:", "ಸಾರ್ವಜನಿಕ ಆಸ್ಪತ್ರೆಗಳಲ್ಲಿ ಟೋಕನ್ ಪಡೆಯಲು ಗಂಟೆಗಟ್ಟಲೆ ಕ್ಯೂ ನಿಲ್ಲುವುದು ವೃದ್ಧರು ಮತ್ತು ರೋಗಿಗಳಿಗೆ ತೀವ್ರ ಕಷ್ಟಕರ."),
    ("೨. ಭಾಷಾ ಸಂವಹನದ ಕೊರತೆ:", "ವಲಸಿಗರು ಮತ್ತು ಹೊರಗಿನಿಂದ ಬಂದವರಿಗೆ ಕನ್ನಡದಲ್ಲಿ ವೈದ್ಯರು, ಆಸ್ಪತ್ರೆ ಸಿಬ್ಬಂದಿ ಹಾಗೂ ಔಷಧಾಲಯಗಳಲ್ಲಿ ಮಾತನಾಡುವುದು ಸವಾಲು."),
    ("೩. ಪ್ರವೇಶಿಸುವಿಕೆಯ ಅಡೆತಡೆಗಳು:", "ಅಂಗವಿಕಲತೆ, ಕೈ ನಡುಕ ಅಥವಾ ಕಣ್ಣಿನ ತೊಂದರೆ ಇರುವವರಿಗೆ ಸಣ್ಣ ಸ್ಕ್ರೀನ್‌ಗಳಲ್ಲಿ ಟೈಪ್ ಮಾಡುವುದು ಅಸಾಧ್ಯ.")
]
for i, (ct, cd) in enumerate(challenges):
    y_pos = Inches(2.7 + i * 1.3)
    c_box = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.1), y_pos, Inches(5.05), Inches(1.15))
    c_box.fill.solid()
    c_box.fill.fore_color.rgb = RGBColor(30, 20, 25)
    c_box.line.color.rgb = RGBColor(185, 28, 28)
    c_box.line.width = Pt(1.0)
    tf_c = c_box.text_frame
    tf_c.word_wrap = True
    p_c1 = tf_c.paragraphs[0]
    p_c1.text = ct
    p_c1.font.name = FONT_NAME
    p_c1.font.size = Pt(13)
    p_c1.font.bold = True
    p_c1.font.color.rgb = RGBColor(254, 202, 202)
    p_c2 = tf_c.add_paragraph()
    p_c2.text = cd
    p_c2.font.name = FONT_NAME
    p_c2.font.size = Pt(11.5)
    p_c2.font.color.rgb = COLOR_MUTED

# Right Card: ಹೀಲಿಯೋ ಪರಿಹಾರ (The Solution)
card2_r = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.85), Inches(5.65), Inches(4.95))
card2_r.fill.solid()
card2_r.fill.fore_color.rgb = COLOR_CARD
card2_r.line.color.rgb = COLOR_TEAL
card2_r.line.width = Pt(2.0)

tx2_rh = slide2.shapes.add_textbox(Inches(7.15), Inches(2.05), Inches(5.0), Inches(0.55))
p2_rh = tx2_rh.text_frame.paragraphs[0]
p2_rh.text = "💡 ಹೀಲಿಯೋ ಪರಿಹಾರ (Our Solution)"
p2_rh.font.name = FONT_NAME
p2_rh.font.size = Pt(20)
p2_rh.font.bold = True
p2_rh.font.color.rgb = COLOR_TEAL_LIGHT

solutions = [
    ("೧. ಸಂಪೂರ್ಣ ಕೈಮುಕ್ತ ವ್ಯವಸ್ಥೆ:", "ಟೈಪಿಂಗ್ ಇಲ್ಲದೆ ಕನ್ನಡ ಧ್ವನಿ ಅಥವಾ ಕಣ್ಣು ನೋಟದಿಂದ ಕೆಲವೇ ಸೆಕೆಂಡುಗಳಲ್ಲಿ ಆಸ್ಪತ್ರೆ ಅಪಾಯಿಂಟ್‌ಮೆಂಟ್ ಬುಕಿಂಗ್."),
    ("೨. ನೂರಕ್ಕೆ ನೂರು ಶುದ್ಧ ಕನ್ನಡ:", "ಇಂಗ್ಲಿಷ್ ಅವಲಂಬನೆಯಿಲ್ಲದೆ ಸ್ಥಳೀಯ ಭಾಷೆಯಲ್ಲೇ ಸಾರ್ವಜನಿಕರಿಗೆ ಸರಳ ಮತ್ತು ಅರ್ಥಗರ್ಭಿತ ಇಂಟರ್ಫೇಸ್."),
    ("೩. ದೈನಂದಿನ ಕನ್ನಡ ಕಲಿಕಾ ಸಂಗಾತಿ:", "ಆಸ್ಪತ್ರೆ, ಬಸ್, ಆಟೋ, ಔಷಧಾಲಯ ಹಾಗೂ ಕಚೇರಿಗಳಲ್ಲಿ ಬಳಸುವ ನೈಜ ಕನ್ನಡ ನುಡಿಗಟ್ಟುಗಳನ್ನು ಆಡಿಯೋ ಸಹಿತ ಕಲಿಸುವ ವ್ಯವಸ್ಥೆ.")
]
for i, (st, sd) in enumerate(solutions):
    y_pos = Inches(2.7 + i * 1.3)
    s_box = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.15), y_pos, Inches(5.05), Inches(1.15))
    s_box.fill.solid()
    s_box.fill.fore_color.rgb = RGBColor(15, 35, 38)
    s_box.line.color.rgb = COLOR_TEAL
    s_box.line.width = Pt(1.0)
    tf_s = s_box.text_frame
    tf_s.word_wrap = True
    p_s1 = tf_s.paragraphs[0]
    p_s1.text = st
    p_s1.font.name = FONT_NAME
    p_s1.font.size = Pt(13)
    p_s1.font.bold = True
    p_s1.font.color.rgb = COLOR_TEAL_LIGHT
    p_s2 = tf_s.add_paragraph()
    p_s2.text = sd
    p_s2.font.name = FONT_NAME
    p_s2.font.size = Pt(11.5)
    p_s2.font.color.rgb = COLOR_MUTED

# ==============================================================================
# SLIDE 3: ಮೂರು ಕೈಮುಕ್ತ ಪ್ರವೇಶಿಸುವಿಕೆ ವಿಧಾನಗಳು
# ==============================================================================
slide3 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(slide3)
add_karnataka_ribbon(slide3)
add_header(slide3, "ಮೂರು ಕೈಮುಕ್ತ ಪ್ರವೇಶಿಸುವಿಕೆ ವಿಧಾನಗಳು", "ವಿಶೇಷ ಚೇತನರು ಹಾಗೂ ಸಾರ್ವಜನಿಕರಿಗೆ ಹೊಂದಿಕೊಳ್ಳುವ ಆಧುನಿಕ ಸಂವಹನ ಮಾಧ್ಯಮಗಳು")
add_footer(slide3, 3)

modes = [
    ("🎙️ ಧ್ವನಿ ವಿಧಾನ (Voice Mode)", "ಕನ್ನಡ ಧ್ವನಿ ಗುರುತಿಸುವಿಕೆ", COLOR_KGOLD, [
        ("ತಂತ್ರಜ್ಞಾನ:", "Web Speech API (kn-IN) ಇಂಟಿಗ್ರೇಷನ್."),
        ("ಬಳಕೆ:", "\"ನನ್ನ ಹೆಸರು ರಮೇಶ್ ಕುಮಾರ್\" ಎಂದು ಮಾತನಾಡಿದರೆ ಸಾಕು."),
        ("ವೈಶಿಷ್ಟ್ಯ:", "ಲೈವ್ ಆಡಿಯೋ ತರಂಗಗಳು ಮತ್ತು ಸ್ಪಷ್ಟ ಕನ್ನಡ ಪಠ್ಯ ಪರಿವರ್ತನೆ."),
        ("ಪ್ರಯೋಜನ:", "ಅನಕ್ಷರಸ್ಥರು ಮತ್ತು ಹಿರಿಯರಿಗೆ ಅತ್ಯಂತ ಸಹಕಾರಿ.")
    ]),
    ("👁️ ಕಣ್ಣು ನೋಟ (Eye Gaze Dwell)", "೧.೨ ಸೆಕೆಂಡ್ SVG ಟೈಮರ್", COLOR_TEAL, [
        ("ತಂತ್ರಜ್ಞಾನ:", "೧.೨ ಸೆಕೆಂಡುಗಳ ವೃತ್ತಾಕಾರದ SVG ಪ್ರೋಗ್ರೆಸ್ ರಿಂಗ್."),
        ("ಬಳಕೆ:", "ಬಟನ್ ಮೇಲೆ ೧.೨ ಸೆಕೆಂಡ್ ದೃಷ್ಟಿ ಅಥವಾ ಪಾಯಿಂಟರ್ ನಿಲ್ಲಿಸಿದಾಗ ಆಟೋ-ಕ್ಲಿಕ್."),
        ("ವೈಶಿಷ್ಟ್ಯ:", "ವೆಬ್ ಆಡಿಯೋ ಚೈಮ್ (Chime) ನೊಂದಿಗೆ ದೃಢೀಕರಣ."),
        ("ಪ್ರಯೋಜನ:", "ಕೈ ಚಲನೆ ಅಸಾಧ್ಯವಾದ ವಿಶೇಷ ಚೇತನರಿಗೆ ವರದಾನ.")
    ]),
    ("👆 ದೊಡ್ಡ ಬಟನ್‌ಗಳು (Touch & Keys)", "WCAG AAA ಅಕ್ಸೆಸಿಬಿಲಿಟಿ", COLOR_KRED, [
        ("ತಂತ್ರಜ್ಞಾನ:", "WCAG AAA ೫೬px+ ಕಾಂಟ್ರಾಸ್ಟ್ ಮತ್ತು ಟಚ್ ಟಾರ್ಗೆಟ್."),
        ("ಬಳಕೆ:", "ಪರದೆಯ ಮೇಲಿನ ದೊಡ್ಡ ಸ್ಪಷ್ಟ ಬಟನ್‌ಗಳನ್ನು ನೇರವಾಗಿ ಮುಟ್ಟಿ."),
        ("ವೈಶಿಷ್ಟ್ಯ:", "ಕೀಬೋರ್ಡ್ Tab, Shift+Tab, Enter, Space ಪೂರ್ಣ ಬೆಂಬಲ."),
        ("ಪ್ರಯೋಜನ:", "ದೃಷ್ಟಿ ಮಾಂದ್ಯತೆ ಮತ್ತು ಸಣ್ಣ ಪರದೆಯ ಮೊಬೈಲ್‌ಗಳಲ್ಲಿ ಸುಲಭ ಬಳಕೆ.")
    ])
]

for i, (m_title, m_tag, border_c, m_bullets) in enumerate(modes):
    x_pos = Inches(0.8 + i * 3.98)
    card_m = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_pos, Inches(1.85), Inches(3.78), Inches(4.95))
    card_m.fill.solid()
    card_m.fill.fore_color.rgb = COLOR_CARD
    card_m.line.color.rgb = border_c
    card_m.line.width = Pt(2.0)

    # Title box
    t_box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_pos + Inches(0.2), Inches(2.05), Inches(3.38), Inches(0.85))
    t_box.fill.solid()
    t_box.fill.fore_color.rgb = RGBColor(15, 23, 42)
    t_box.line.color.rgb = border_c
    t_box.line.width = Pt(1.0)
    tf_mb = t_box.text_frame
    p_mb1 = tf_mb.paragraphs[0]
    p_mb1.alignment = PP_ALIGN.CENTER
    p_mb1.text = m_title
    p_mb1.font.name = FONT_NAME
    p_mb1.font.size = Pt(14)
    p_mb1.font.bold = True
    p_mb1.font.color.rgb = COLOR_WHITE
    p_mb2 = tf_mb.add_paragraph()
    p_mb2.alignment = PP_ALIGN.CENTER
    p_mb2.text = m_tag
    p_mb2.font.name = FONT_NAME
    p_mb2.font.size = Pt(11)
    p_mb2.font.bold = True
    p_mb2.font.color.rgb = border_c

    # Bullets
    tx_b = slide3.shapes.add_textbox(x_pos + Inches(0.2), Inches(3.05), Inches(3.38), Inches(3.55))
    tf_b = tx_b.text_frame
    tf_b.word_wrap = True
    for j, (lbl, desc) in enumerate(m_bullets):
        p_b = tf_b.paragraphs[0] if j == 0 else tf_b.add_paragraph()
        p_b.space_after = Pt(10)
        run1 = p_b.add_run()
        run1.text = f"• {lbl} "
        run1.font.name = FONT_NAME
        run1.font.size = Pt(12)
        run1.font.bold = True
        run1.font.color.rgb = border_c
        run2 = p_b.add_run()
        run2.text = desc
        run2.font.name = FONT_NAME
        run2.font.size = Pt(11.5)
        run2.font.color.rgb = COLOR_MUTED

# ==============================================================================
# SLIDE 4: ರೋಗಿಯ ಗುರುತು & ೫ ಪ್ರಮುಖ ವೈದ್ಯಕೀಯ ವಿಭಾಗಗಳು
# ==============================================================================
slide4 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(slide4)
add_karnataka_ribbon(slide4)
add_header(slide4, "ರೋಗಿಯ ಗುರುತು ದಾಖಲಾತಿ ಮತ್ತು ೫ ವೈದ್ಯಕೀಯ ವಿಭಾಗಗಳು", "ಕನ್ನಡ NLP ಮೂಲಕ ಹೆಸರು ಗುರುತಿಸುವಿಕೆ ಮತ್ತು ತಜ್ಞ ವೈದ್ಯರ ಪಟ್ಟಿ")
add_footer(slide4, 4)

# Left: Voice Identity
card4_l = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.85), Inches(5.65), Inches(4.95))
card4_l.fill.solid()
card4_l.fill.fore_color.rgb = COLOR_CARD
card4_l.line.color.rgb = COLOR_KGOLD
card4_l.line.width = Pt(2.0)

tx4_lh = slide4.shapes.add_textbox(Inches(1.1), Inches(2.05), Inches(5.0), Inches(0.55))
p4_lh = tx4_lh.text_frame.paragraphs[0]
p4_lh.text = "🎙️ ಧ್ವನಿ ಮೂಲಕ ರೋಗಿಯ ಗುರುತು ದಾಖಲು"
p4_lh.font.name = FONT_NAME
p4_lh.font.size = Pt(18)
p4_lh.font.bold = True
p4_lh.font.color.rgb = COLOR_KGOLD

v_items = [
    ("ಧ್ವನಿ ಹೆಸರು ದಾಖಲಾತಿ:", "ರೋಗಿಯು 'ನನ್ನ ಹೆಸರು ರಮೇಶ್ ಕುಮಾರ್' ಅಥವಾ 'ಅನಿತಾ ಗೌಡ' ಎಂದು ಮಾತನಾಡಿದಾಗ ಕನ್ನಡ NLP ತಕ್ಷಣವೇ ಗುರುತಿಸಿ ಇನ್‌ಪುಟ್ ಬಾಕ್ಸ್‌ನಲ್ಲಿ ನಮೂದಿಸುತ್ತದೆ."),
    ("ಹುಟ್ಟಿದ ದಿನಾಂಕ (DOB):", "ಧ್ವನಿ ಮೂಲಕ ಹುಟ್ಟಿದ ದಿನಾಂಕ ಅಥವಾ ವರ್ಷವನ್ನು ಮಾತನಾಡಬಹುದು (ಉದಾ: '೧೫/೦೬/೧೯೮೮')."),
    ("೪-ಅಂಕಿಯ ಡೆಮೊ ಕೋಡ್:", "ಪ್ರಾತ್ಯಕ್ಷಿಕೆಯಲ್ಲಿ ಅಧಿಕೃತ ದೃಢೀಕರಣಕ್ಕಾಗಿ ಕನ್ನಡ ಅಂಕಿಗಳಲ್ಲಿ '೪೨೮೫' ಕೋಡ್ ಪ್ರದರ್ಶನ."),
    ("ಗೌಪ್ಯತೆ & ಸುರಕ್ಷತೆ:", "ಯಾವುದೇ ಧ್ವನಿ ರೆಕಾರ್ಡ್ ಆಗುವುದಿಲ್ಲ. ಮಾಹಿತಿ ಸಾಧನದಿಂದ ಹೊರಹೋಗುವುದಿಲ್ಲ.")
]
for i, (vt, vd) in enumerate(v_items):
    y_pos = Inches(2.65 + i * 1.0)
    vb = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.1), y_pos, Inches(5.05), Inches(0.9))
    vb.fill.solid()
    vb.fill.fore_color.rgb = RGBColor(15, 23, 42)
    vb.line.color.rgb = RGBColor(71, 85, 105)
    vb.line.width = Pt(1.0)
    tf_vb = vb.text_frame
    tf_vb.word_wrap = True
    p_vb1 = tf_vb.paragraphs[0]
    p_vb1.text = vt
    p_vb1.font.name = FONT_NAME
    p_vb1.font.size = Pt(12)
    p_vb1.font.bold = True
    p_vb1.font.color.rgb = COLOR_KGOLD_LIGHT
    p_vb2 = tf_vb.add_paragraph()
    p_vb2.text = vd
    p_vb2.font.name = FONT_NAME
    p_vb2.font.size = Pt(10.5)
    p_vb2.font.color.rgb = COLOR_MUTED

# Right: 5 Departments
card4_r = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.85), Inches(5.65), Inches(4.95))
card4_r.fill.solid()
card4_r.fill.fore_color.rgb = COLOR_CARD
card4_r.line.color.rgb = COLOR_TEAL
card4_r.line.width = Pt(2.0)

tx4_rh = slide4.shapes.add_textbox(Inches(7.15), Inches(2.05), Inches(5.0), Inches(0.55))
p4_rh = tx4_rh.text_frame.paragraphs[0]
p4_rh.text = "🏥 ೫ ಪ್ರಮುಖ ಆಸ್ಪತ್ರೆ ವಿಭಾಗಗಳು & ತಜ್ಞರು"
p4_rh.font.name = FONT_NAME
p4_rh.font.size = Pt(18)
p4_rh.font.bold = True
p4_rh.font.color.rgb = COLOR_TEAL_LIGHT

depts = [
    ("❤️ ಹೃದಯ ವಿಭಾಗ (Cardiology)", "ಡಾ. ರಮೇಶ್ ಕುಮಾರ್", "ಎದೆ ನೋವು, ರಕ್ತದೊತ್ತಡ ಮತ್ತು ಹೃದಯ ತಪಾಸಣೆ"),
    ("👶 ಮಕ್ಕಳ ವಿಭಾಗ (Paediatrics)", "ಡಾ. ಅನಿತಾ ರಾವ್", "ಮಕ್ಕಳ ಜ್ವರ, ರೋಗನಿರೋಧಕ ಲಸಿಕೆ ಮತ್ತು ಆರೈಕೆ"),
    ("🦴 ಮೂಳೆ ವಿಭಾಗ (Orthopaedics)", "ಡಾ. ಸುರೇಶ್ ಗೌಡ", "ಕೀಲು ನೋವು, ಮೂಳೆ ಮುರಿತ ಮತ್ತು ಸಂಧಿವಾತ"),
    ("👁️ ಕಣ್ಣು ವಿಭಾಗ (Ophthalmology)", "ಡಾ. ಪ್ರಿಯಾ ಶರ್ಮ", "ದೃಷ್ಟಿ ಪರೀಕ್ಷೆ, ಕಣ್ಣಿನ ಪೊರೆ ಮತ್ತು ಕನ್ನಡಕ"),
    ("🩺 ಸಾಮಾನ್ಯ ವೈದ್ಯ (General Medicine)", "ಡಾ. ಮಂಜುನಾಥ್ ಕೆ.", "ಜ್ವರ, ಶೀತ, ಮಧುಮೇಹ ಮತ್ತು ಸಾಮಾನ್ಯ ತಪಾಸಣೆ")
]
for i, (d_name, doc, desc) in enumerate(depts):
    y_pos = Inches(2.65 + i * 0.8)
    db = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.15), y_pos, Inches(5.05), Inches(0.72))
    db.fill.solid()
    db.fill.fore_color.rgb = RGBColor(15, 23, 42)
    db.line.color.rgb = COLOR_TEAL
    db.line.width = Pt(1.0)
    tf_db = db.text_frame
    tf_db.word_wrap = True
    p_db1 = tf_db.paragraphs[0]
    p_db1.text = f"{d_name} — {doc}"
    p_db1.font.name = FONT_NAME
    p_db1.font.size = Pt(12)
    p_db1.font.bold = True
    p_db1.font.color.rgb = COLOR_TEAL_LIGHT
    p_db2 = tf_db.add_paragraph()
    p_db2.text = desc
    p_db2.font.name = FONT_NAME
    p_db2.font.size = Pt(10.5)
    p_db2.font.color.rgb = COLOR_MUTED

# ==============================================================================
# SLIDE 5: ನೈಜ-ಸಮಯದ ಸ್ಲಾಟ್ ಬುಕಿಂಗ್ & ಮೊಬೈಲ್ ಸಂದೇಶ
# ==============================================================================
slide5 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(slide5)
add_karnataka_ribbon(slide5)
add_header(slide5, "ನೈಜ-ಸಮಯದ ಸ್ಲಾಟ್ ಬುಕಿಂಗ್ & SMS ದೃಢೀಕರಣ", "ಕ್ಯಾಲೆಂಡರ್ ವೇಳಾಪಟ್ಟಿ, ಅಧಿಕೃತ ಉಲ್ಲೇಖ ಸಂಕೇತ ಮತ್ತು ಸಕ್ರಿಯ ಬುಕಿಂಗ್ ನಿರ್ವಹಣೆ")
add_footer(slide5, 5)

# Left: Real-time Slot Scheduling
card5_l = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.85), Inches(5.65), Inches(4.95))
card5_l.fill.solid()
card5_l.fill.fore_color.rgb = COLOR_CARD
card5_l.line.color.rgb = COLOR_TEAL
card5_l.line.width = Pt(2.0)

tx5_lh = slide5.shapes.add_textbox(Inches(1.1), Inches(2.05), Inches(5.0), Inches(0.55))
p5_lh = tx5_lh.text_frame.paragraphs[0]
p5_lh.text = "📅 ಕ್ಯಾಲೆಂಡರ್ & ನೈಜ-ಸಮಯದ ಸ್ಲಾಟ್‌ಗಳು"
p5_lh.font.name = FONT_NAME
p5_lh.font.size = Pt(18)
p5_lh.font.bold = True
p5_lh.font.color.rgb = COLOR_TEAL_LIGHT

slot_features = [
    ("ಮುಂದಿನ ೭ ದಿನಗಳ ಕ್ಯಾಲೆಂಡರ್:", "ಇಂದು (ಸೋಮವಾರ), ನಾಳೆ (ಮಂಗಳವಾರ), ಬುಧವಾರ... ಸುಲಭವಾಗಿ ದಿನಾಂಕ ಆಯ್ಕೆ."),
    ("ಬೆಳಿಗ್ಗೆ & ಮಧ್ಯಾಹ್ನದ ಸ್ಲಾಟ್‌ಗಳು:", "೧೦:೦೦, ೧೦:೩೦, ೧೧:೦೦ ಬೆಳಿಗ್ಗೆ ಮತ್ತು ೦೨:೦೦, ೦೨:೩೦, ೦೩:೦೦ ಮಧ್ಯಾಹ್ನ."),
    ("ನೈಜ-ಸಮಯದ REST API ಸಿಂಕ್:", "ಯಾರಾದರೂ ಸ್ಲಾಟ್ ಬುಕ್ ಮಾಡಿದ ತಕ್ಷಣ ಲಭ್ಯತೆ ಕೆಂಪು ಬಣ್ಣಕ್ಕೆ ತಿರುಗುತ್ತದೆ. ಲಭ್ಯವಿರುವ ಸ್ಲಾಟ್‌ಗಳು ಹಸಿರಾಗಿರುತ್ತವೆ."),
    ("ಡೇಟಾಬೇಸ್ ಸಂರಕ್ಷಣೆ:", "Node.js ಬ್ಯಾಕೆಂಡ್ ಮೂಲಕ `data/availability.json` ನಲ್ಲಿ ನೈಜ-ಸಮಯದ ಆಟೋಮ್ಯಾಟಿಕ್ ರಿಸರ್ವೇಶನ್.")
]
for i, (sft, sfd) in enumerate(slot_features):
    y_pos = Inches(2.65 + i * 1.0)
    sfb = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.1), y_pos, Inches(5.05), Inches(0.9))
    sfb.fill.solid()
    sfb.fill.fore_color.rgb = RGBColor(15, 23, 42)
    sfb.line.color.rgb = COLOR_TEAL
    sfb.line.width = Pt(1.0)
    tf_sfb = sfb.text_frame
    tf_sfb.word_wrap = True
    p_sfb1 = tf_sfb.paragraphs[0]
    p_sfb1.text = sft
    p_sfb1.font.name = FONT_NAME
    p_sfb1.font.size = Pt(12)
    p_sfb1.font.bold = True
    p_sfb1.font.color.rgb = COLOR_TEAL_LIGHT
    p_sfb2 = tf_sfb.add_paragraph()
    p_sfb2.text = sfd
    p_sfb2.font.name = FONT_NAME
    p_sfb2.font.size = Pt(10.5)
    p_sfb2.font.color.rgb = COLOR_MUTED

# Right: Confirmation & Management
card5_r = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.85), Inches(5.65), Inches(4.95))
card5_r.fill.solid()
card5_r.fill.fore_color.rgb = COLOR_CARD
card5_r.line.color.rgb = COLOR_KGOLD
card5_r.line.width = Pt(2.0)

tx5_rh = slide5.shapes.add_textbox(Inches(7.15), Inches(2.05), Inches(5.0), Inches(0.55))
p5_rh = tx5_rh.text_frame.paragraphs[0]
p5_rh.text = "🎉 ದೃಢೀಕರಣ, SMS & ಬುಕಿಂಗ್ ನಿರ್ವಹಣೆ"
p5_rh.font.name = FONT_NAME
p5_rh.font.size = Pt(18)
p5_rh.font.bold = True
p5_rh.font.color.rgb = COLOR_KGOLD

conf_features = [
    ("ಅಧಿಕೃತ ಕನ್ನಡ ಉಲ್ಲೇಖ ಸಂಖ್ಯೆ:", "ಬುಕಿಂಗ್ ಮುಗಿದ ತಕ್ಷಣ ಕನ್ನಡ ಅಂಕಿಗಳಲ್ಲಿ `HL-೪೨೮೫` ವಿಶಿಷ್ಟ ರೆಫರೆನ್ಸ್ ಕೋಡ್ ಸೃಷ್ಟಿಯಾಗುತ್ತದೆ."),
    ("ಮೊಬೈಲ್ SMS ಪ್ರಿವ್ಯೂ:", "\"ಆಸ್ಪತ್ರೆ ಸಂದೇಶ: ಗೌರವಾನ್ವಿತ ರಮೇಶ್ ಕುಮಾರ್, ನಿಮ್ಮ ಅಪಾಯಿಂಟ್‌ಮೆಂಟ್ ಯಶಸ್ವಿಯಾಗಿದೆ. ಉಲ್ಲೇಖ: HL-೪೨೮೫\"."),
    ("ನನ್ನ ಬುಕಿಂಗ್‌ಗಳು ನಿರ್ವಹಣೆ:", "ಸಕ್ರಿಯ ಅಪಾಯಿಂಟ್‌ಮೆಂಟ್‌ಗಳ ಪಟ್ಟಿ, ದಿನಾಂಕ ಮರುಹೊಂದಿಸುವ (Reschedule) ಸೌಲಭ್ಯ."),
    ("ಸ್ವಯಂಚಾಲಿತ ಸ್ಲಾಟ್ ಬಿಡುಗಡೆ:", "ಅಪಾಯಿಂಟ್‌ಮೆಂಟ್ ರದ್ದು (Cancel) ಮಾಡಿದಾಗ ಬ್ಯಾಕೆಂಡ್‌ನಲ್ಲಿ ಆ ಸ್ಲಾಟ್ ಸ್ವಯಂಚಾಲಿತವಾಗಿ ಇತರ ರೋಗಿಗಳಿಗೆ ಮುಕ್ತವಾಗುತ್ತದೆ.")
]
for i, (cft, cfd) in enumerate(conf_features):
    y_pos = Inches(2.65 + i * 1.0)
    cfb = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.15), y_pos, Inches(5.05), Inches(0.9))
    cfb.fill.solid()
    cfb.fill.fore_color.rgb = RGBColor(15, 23, 42)
    cfb.line.color.rgb = COLOR_KGOLD
    cfb.line.width = Pt(1.0)
    tf_cfb = cfb.text_frame
    tf_cfb.word_wrap = True
    p_cfb1 = tf_cfb.paragraphs[0]
    p_cfb1.text = cft
    p_cfb1.font.name = FONT_NAME
    p_cfb1.font.size = Pt(12)
    p_cfb1.font.bold = True
    p_cfb1.font.color.rgb = COLOR_KGOLD_LIGHT
    p_cfb2 = tf_cfb.add_paragraph()
    p_cfb2.text = cfd
    p_cfb2.font.name = FONT_NAME
    p_cfb2.font.size = Pt(10.5)
    p_cfb2.font.color.rgb = COLOR_MUTED

# ==============================================================================
# SLIDE 6: ದೈನಂದಿನ ಜೀವನದಲ್ಲಿ ಕನ್ನಡ ಕಲಿಕೆ & ೧೦೮ ತುರ್ತು ಪ್ರೋಟೋಕಾಲ್
# ==============================================================================
slide6 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(slide6)
add_karnataka_ribbon(slide6)
add_header(slide6, "ದೈನಂದಿನ ಜೀವನದಲ್ಲಿ ಕನ್ನಡ ಕಲಿಕೆ & ೧೦೮ ತುರ್ತು ಪ್ರೋಟೋಕಾಲ್", "ನೈಜ ಭಾಷಾ ತರಬೇತಿ ಸಂಗಾತಿ ಮತ್ತು ಜೀವ ರಕ್ಷಣಾ ಸುರಕ್ಷತಾ ತಂತ್ರಜ್ಞಾನ")
add_footer(slide6, 6)

# Left: Learn Kannada
card6_l = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.85), Inches(5.65), Inches(4.95))
card6_l.fill.solid()
card6_l.fill.fore_color.rgb = COLOR_CARD
card6_l.line.color.rgb = COLOR_KGOLD
card6_l.line.width = Pt(2.0)

tx6_lh = slide6.shapes.add_textbox(Inches(1.1), Inches(2.05), Inches(5.0), Inches(0.55))
p6_lh = tx6_lh.text_frame.paragraphs[0]
p6_lh.text = "🎓 ದೈನಂದಿನ ಜೀವನದಲ್ಲಿ ಕನ್ನಡ (Learning Center)"
p6_lh.font.name = FONT_NAME
p6_lh.font.size = Pt(18)
p6_lh.font.bold = True
p6_lh.font.color.rgb = COLOR_KGOLD

learn_points = [
    ("೫ ನೈಜ ಜೀವನದ ವಿಭಾಗಗಳು:", "🏥 ಆರೋಗ್ಯ  |  🚌 ಪ್ರಯಾಣ & ಸಾರಿಗೆ  |  💊 ಔಷಧಾಲಯ & ಶಾಪಿಂಗ್  |  🏢 ಕಚೇರಿ  |  🚨 ತುರ್ತು ಪರಿಸ್ಥಿತಿ."),
    ("ಪ್ರಾಯೋಗಿಕ ನುಡಿಗಟ್ಟುಗಳು:", "\"ಆಸ್ಪತ್ರೆಗೆ ಹೋಗಲು ಬಸ್ ಯಾವುದು?\", \"ಈ ಮಾತ್ರೆಗಳು ಲಭ್ಯವಿದೆಯೇ?\", \"ಆನ್‌ಲೈನ್ ಪಾವತಿ (UPI) ಮಾಡಬಹುದೇ?\" ಮುಂತಾದವು."),
    ("ಆಡಿಯೋ ಉಚ್ಚಾರಣೆ (Listen):", "ಪ್ರತಿಯೊಂದು ವಾಕ್ಯಕ್ಕೂ ಸ್ಪಷ್ಟ ಕನ್ನಡ ಧ್ವನಿ, ಇಂಗ್ಲಿಷ್ ಅರ್ಥ ಮತ್ತು ನೈಜ ಸಂದರ್ಭದ ವಿವರಣೆ."),
    ("ಲೈವ್ ಧ್ವನಿ ಅಭ್ಯಾಸ ಪರೀಕ್ಷೆ:", "ಬಳಕೆದಾರರು ಮಾತನಾಡಿ ಉಚ್ಚಾರಣೆ ಪರೀಕ್ಷಿಸಬಹುದು — 'ಶೇಕಡಾ ೯೬ ಹೊಂದಾಣಿಕೆ (Excellent Match)' ಫಲಿತಾಂಶ!")
]
for i, (lt, ld) in enumerate(learn_points):
    y_pos = Inches(2.65 + i * 1.0)
    lb = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.1), y_pos, Inches(5.05), Inches(0.9))
    lb.fill.solid()
    lb.fill.fore_color.rgb = RGBColor(15, 23, 42)
    lb.line.color.rgb = COLOR_KGOLD
    lb.line.width = Pt(1.0)
    tf_lb = lb.text_frame
    tf_lb.word_wrap = True
    p_lb1 = tf_lb.paragraphs[0]
    p_lb1.text = lt
    p_lb1.font.name = FONT_NAME
    p_lb1.font.size = Pt(12)
    p_lb1.font.bold = True
    p_lb1.font.color.rgb = COLOR_KGOLD_LIGHT
    p_lb2 = tf_lb.add_paragraph()
    p_lb2.text = ld
    p_lb2.font.name = FONT_NAME
    p_lb2.font.size = Pt(10.5)
    p_lb2.font.color.rgb = COLOR_MUTED

# Right: Emergency 108
card6_r = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.85), Inches(5.65), Inches(4.95))
card6_r.fill.solid()
card6_r.fill.fore_color.rgb = COLOR_CARD
card6_r.line.color.rgb = COLOR_KRED
card6_r.line.width = Pt(2.0)

tx6_rh = slide6.shapes.add_textbox(Inches(7.15), Inches(2.05), Inches(5.0), Inches(0.55))
p6_rh = tx6_rh.text_frame.paragraphs[0]
p6_rh.text = "🚨 ತುರ್ತು ೧೦೮ ಸುರಕ್ಷತಾ ಪ್ರೋಟೋಕಾಲ್"
p6_rh.font.name = FONT_NAME
p6_rh.font.size = Pt(18)
p6_rh.font.bold = True
p6_rh.font.color.rgb = COLOR_KRED

em_points = [
    ("ತುರ್ತು ಕೀವರ್ಡ್ ಪತ್ತೆ:", "ರೋಗಿ ಮಾತಿನಲ್ಲಿ 'ಎದೆ ನೋವು', 'ಉಸಿರಾಟದ ತೊಂದರೆ', 'ರಕ್ತಸ್ರಾವ', 'ಹೃದಯಾಘಾತ', 'ಪ್ರಜ್ಞೆ ತಪ್ಪಿದೆ' ಪದಗಳು ಕೇಳಿಬಂದಾಗ NLP ತಕ್ಷಣ ಎಚ್ಚೆತ್ತುಕೊಳ್ಳುತ್ತದೆ."),
    ("ಬುಕಿಂಗ್ ಪ್ರಕ್ರಿಯೆ ಸ್ಥಗಿತ:", "ಆನ್‌ಲೈನ್ ಅಪಾಯಿಂಟ್‌ಮೆಂಟ್ ಬುಕಿಂಗ್ ಅನ್ನು ತಕ್ಷಣ ನಿಲ್ಲಿಸಿ ಜೀವ ರಕ್ಷಣೆಗೆ ಮೊದಲ ಆದ್ಯತೆ ನೀಡಲಾಗುತ್ತದೆ."),
    ("ಪೂರ್ಣ ಪರದೆಯ ತುರ್ತು ಎಚ್ಚರಿಕೆ:", "ಕೆಂಪು ಬಣ್ಣದ ತುರ್ತು ಪರದೆ ಮತ್ತು ಎಚ್ಚರಿಕೆಯ ಚೈಮ್ ಸೌಂಡ್ ತಕ್ಷಣ ಪ್ರದರ್ಶನಗೊಳ್ಳುತ್ತದೆ."),
    ("ನೇರ ೧೦೮ ಆಂಬ್ಯುಲೆನ್ಸ್ ಡಯಲರ್:", "'📞 ೧೦೮ ಗೆ ಕರೆ ಮಾಡಿ' ಬಟನ್ ಮೂಲಕ ಉಚಿತ ಸರ್ಕಾರಿ ಆಂಬ್ಯುಲೆನ್ಸ್ ಸೇವೆಗೆ ನೇರ ಫೋನ್ ಸಂಪರ್ಕ.")
]
for i, (et, ed) in enumerate(em_points):
    y_pos = Inches(2.65 + i * 1.0)
    eb = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.15), y_pos, Inches(5.05), Inches(0.9))
    eb.fill.solid()
    eb.fill.fore_color.rgb = RGBColor(30, 20, 25)
    eb.line.color.rgb = COLOR_KRED
    eb.line.width = Pt(1.0)
    tf_eb = eb.text_frame
    tf_eb.word_wrap = True
    p_eb1 = tf_eb.paragraphs[0]
    p_eb1.text = et
    p_eb1.font.name = FONT_NAME
    p_eb1.font.size = Pt(12)
    p_eb1.font.bold = True
    p_eb1.font.color.rgb = RGBColor(254, 202, 202)
    p_eb2 = tf_eb.add_paragraph()
    p_eb2.text = ed
    p_eb2.font.name = FONT_NAME
    p_eb2.font.size = Pt(10.5)
    p_eb2.font.color.rgb = COLOR_MUTED

# ==============================================================================
# SLIDE 7: ತಾಂತ್ರಿಕ ವಾಸ್ತುಶಿಲ್ಪ & ಸಮಾರೋಪ
# ==============================================================================
slide7 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(slide7)
add_karnataka_ribbon(slide7)
add_header(slide7, "ತಾಂತ್ರಿಕ ವಾಸ್ತುಶಿಲ್ಪ, ಪರೀಕ್ಷೆಗಳು & ಸಮಾರೋಪ", "ಶೂನ್ಯ-ಅವಲಂಬನೆ Node.js ಬ್ಯಾಕೆಂಡ್, WebGL ೩ಡಿ ಹಾಗೂ ೧೦೦% ಪರೀಕ್ಷಾ ಕವರೇಜ್")
add_footer(slide7, 7)

# 3 Tech Columns
tech_cols = [
    ("💻 ಶೂನ್ಯ-ಅವಲಂಬನೆ REST ಬ್ಯಾಕೆಂಡ್", "Node.js Architecture", COLOR_TEAL, [
        ("ಸ್ಥಳೀಯ Node.js:", "ಯಾವುದೇ ಬಾಹ್ಯ npm ಲೈಬ್ರರಿಗಳಿಲ್ಲದೆ ಶುದ್ಧ Node.js (http, fs, path, url)."),
        ("REST API ಎಂಡ್‌ಪಾಯಿಂಟ್‌ಗಳು:", "GET /api/departments, GET /api/slots, POST /api/bookings, DELETE /api/bookings/:id."),
        ("ಪರ್ಸಿಸ್ಟೆಂಟ್ JSON DB:", "bookings.json, availability.json, learn_phrases.json."),
        ("ಸ್ಥಳೀಯ ಫಾಲ್‌ಬ್ಯಾಕ್:", "ಸರ್ವರ್ ಇಲ್ಲದಿದ್ದರೂ ಬ್ರೌಸರ್ localStorage ಮೂಲಕ ೧೦೦% ಆಫ್‌ಲೈನ್ ರನ್.")
    ]),
    ("🧪 ೩೫ ಸ್ವಯಂಚಾಲಿತ NLP ಪರೀಕ್ಷೆಗಳು", "100% Automated Testing", COLOR_GREEN, [
        ("ಟೆಸ್ಟ್ ಸೂಟ್ (tests.html):", "ಕನ್ನಡ ನೈಸರ್ಗಿಕ ಭಾಷಾ ಸಂಸ್ಕರಣೆಯ ೩೫ ಯೂನಿಟ್ ಪರೀಕ್ಷೆಗಳು."),
        ("ಇಂಟೆಂಟ್‌ಗಳ ಪರೀಕ್ಷೆ:", "book, cancel, reschedule, yes, no, help, emergency ೧೦೦% Pass."),
        ("ಕನ್ನಡ ಅಂಕಿ & ದಿನಾಂಕ:", "ಇಂದು, ನಾಳೆ, ದಿನಗಳು ಹಾಗೂ ಕನ್ನಡ ಅಂಕಿಗಳ ಸಮರ್ಥ ಪಾರ್ಸಿಂಗ್."),
        ("ಟೈಪೋ ಸಹಿಷ್ಣುತೆ:", "ಲೆವೆನ್‌ಸ್ಟೈನ್ ದೂರ ಅಲ್ಗಾರಿದಮ್ ಮೂಲಕ ಅಕ್ಷರ ತಪ್ಪುಗಳಿದ್ದರೂ ನಿಖರ ಗ್ರಹಿಕೆ.")
    ]),
    ("🌐 WebGL & Three.js ೩ಡಿ ಎಂಜಿನ್", "Interactive 3D Graphics", COLOR_KGOLD, [
        ("Three.js r128 ಪವರ್ಡ್:", "ಹಾರ್ಡ್‌ವೇರ್ ಆಕ್ಸಲರೇಟೆಡ್ ೩ಡಿ ರೌಂಡೆಡ್ ಮೆಡಿಕಲ್ ಕ್ರಾಸ್."),
        ("ಶಾಂತ ಬೆಳಕಿನ ಸಂಯೋಜನೆ:", "ಟೀಲ್/ಸಯಾನ್ ಫೋನ್ ಮೆಟೀರಿಯಲ್ & ೩ಡಿ ಹ್ಯಾಲೋ ರಿಂಗ್."),
        ("ಚಲನೆ ನಿಯಂತ್ರಣ:", "ಬಳಕೆದಾರರಿಗೆ ಚಲನೆ ನಿಲ್ಲಿಸಲು '⏸️ ಚಲನೆ ನಿಲ್ಲಿಸಿ' ಸ್ವಿಚ್ ಸೌಲಭ್ಯ."),
        ("ಸ್ವಯಂಚಾಲಿತ ಫಾಲ್‌ಬ್ಯಾಕ್:", "WebGL ಲಭ್ಯವಿಲ್ಲದಿದ್ದಲ್ಲಿ ಸುಂದರ CSS ಗ್ರೇಡಿಯಂಟ್ ಬ್ಯಾಕ್‌ಡ್ರಾಪ್.")
    ])
]

for i, (t_title, t_tag, t_color, t_bullets) in enumerate(tech_cols):
    x_pos = Inches(0.8 + i * 3.98)
    card_t = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_pos, Inches(1.85), Inches(3.78), Inches(3.55))
    card_t.fill.solid()
    card_t.fill.fore_color.rgb = COLOR_CARD
    card_t.line.color.rgb = t_color
    card_t.line.width = Pt(2.0)

    # Title box
    tb = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_pos + Inches(0.2), Inches(2.0), Inches(3.38), Inches(0.7))
    tb.fill.solid()
    tb.fill.fore_color.rgb = RGBColor(15, 23, 42)
    tb.line.color.rgb = t_color
    tb.line.width = Pt(1.0)
    tf_tb = tb.text_frame
    p_tb1 = tf_tb.paragraphs[0]
    p_tb1.alignment = PP_ALIGN.CENTER
    p_tb1.text = t_title
    p_tb1.font.name = FONT_NAME
    p_tb1.font.size = Pt(13)
    p_tb1.font.bold = True
    p_tb1.font.color.rgb = COLOR_WHITE
    p_tb2 = tf_tb.add_paragraph()
    p_tb2.alignment = PP_ALIGN.CENTER
    p_tb2.text = t_tag
    p_tb2.font.name = FONT_NAME
    p_tb2.font.size = Pt(10.5)
    p_tb2.font.bold = True
    p_tb2.font.color.rgb = t_color

    # Bullets
    tx_tb = slide7.shapes.add_textbox(x_pos + Inches(0.2), Inches(2.8), Inches(3.38), Inches(2.5))
    tf_tb2 = tx_tb.text_frame
    tf_tb2.word_wrap = True
    for j, (lbl, desc) in enumerate(t_bullets):
        p_b = tf_tb2.paragraphs[0] if j == 0 else tf_tb2.add_paragraph()
        p_b.space_after = Pt(4)
        run1 = p_b.add_run()
        run1.text = f"• {lbl} "
        run1.font.name = FONT_NAME
        run1.font.size = Pt(10.5)
        run1.font.bold = True
        run1.font.color.rgb = t_color
        run2 = p_b.add_run()
        run2.text = desc
        run2.font.name = FONT_NAME
        run2.font.size = Pt(10)
        run2.font.color.rgb = COLOR_MUTED

# Grand Outro Banner at Bottom
banner = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.6), Inches(11.733), Inches(1.25))
banner.fill.solid()
banner.fill.fore_color.rgb = RGBColor(185, 28, 28) # Karnataka Red
banner.line.color.rgb = COLOR_KGOLD
banner.line.width = Pt(2.5)

tf_ban = banner.text_frame
tf_ban.word_wrap = True
p_ban1 = tf_ban.paragraphs[0]
p_ban1.alignment = PP_ALIGN.CENTER
p_ban1.text = "👑 ಸಿರಿಗನ್ನಡಂ ಗೆಲ್ಗೆ, ಸಿರಿಗನ್ನಡಂ ಬಾಳ್ಗೆ! 👑"
p_ban1.font.name = FONT_NAME
p_ban1.font.size = Pt(22)
p_ban1.font.bold = True
p_ban1.font.color.rgb = COLOR_KGOLD_LIGHT

p_ban2 = tf_ban.add_paragraph()
p_ban2.alignment = PP_ALIGN.CENTER
p_ban2.text = "ಹೀಲಿಯೋ — ತಂತ್ರಜ್ಞಾನದಲ್ಲಿ ಕನ್ನಡದ ಸಾರ್ಥಕ ಹೆಜ್ಜೆ • ಧನ್ಯವಾದಗಳು!"
p_ban2.font.name = FONT_NAME
p_ban2.font.size = Pt(15)
p_ban2.font.bold = True
p_ban2.font.color.rgb = COLOR_WHITE

# Save Presentation
out_pptx = "d:/kannada/Kannada_Everyday_Life_HEALIO.pptx"
prs.save(out_pptx)
print(f"Presentation saved successfully to: {out_pptx}")
print(f"Total Slides: {len(prs.slides)}")
